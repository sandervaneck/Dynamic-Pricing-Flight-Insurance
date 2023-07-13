import pandas as pd
import statsmodels.api as sm
from openpyxl import Workbook
from openpyxl.utils.dataframe import dataframe_to_rows
import matplotlib.pyplot as plt

def model_probit(train_X, train_Y, test_X, workbook, file):
    sheet = workbook.create_sheet('Probit')
    X = sm.add_constant(train_X)  # Add a constant term to the design matrix
    Y = train_Y
    X_test = sm.add_constant(test_X)

    # Create the probit model
    probit_model = sm.Probit(Y, X)

    # Fit the probit model to the data
    probit_results = probit_model.fit()

    summary_df = pd.DataFrame(probit_results.summary().tables[1].data[1:], columns=probit_results.summary().tables[1].data[0])

    # Write the summary table to the Excel sheet
    for row in dataframe_to_rows(summary_df, index=False, header=True):
        sheet.append(row)

    # Predict the values of 'refund' using the trained model
    probabilities = probit_results.predict(X_test)
    workbook.save(file)

    return probabilities